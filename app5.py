"""
Excel Template Maker — Shareholder Lists Generator
====================================================
What this program does (overview):
1. Logs into aksjeeiere.no via Selenium and downloads xlsx files with each
   company's shareholder list for every year (2014 -> current year - 1).
2. Processes all downloaded xlsx files, merges them into a single Excel
   report based on the Template(original).xlsx template, and calculates
   year-over-year changes in share ownership.
3. (Optional) adds phone numbers via a separate script, telephone.py.
4. Displays everything through a simple Streamlit interface.

app3.py version notes:
- Fixed a bug that crashed on int(' ') in the SEC.NR. / ORG.NR. column.
- ChromeDriver is now selected automatically (webdriver-manager) + headless
  mode, which is required to run on a server.
- Added more checks and debug messages at "fragile" points in the code
  (empty files, missing columns, non-numeric values, etc.).
- Code split into sections with comments for easier navigation.
- CONCURRENCY FIX: every run now uses its own isolated session folder
  (sessions/<session_id>/) instead of one shared "files1" folder for
  everyone. This means several people can use the app on the server at the
  same time without their downloaded files, reports, or cleanup steps
  interfering with each other's data. Chrome now also downloads directly
  into that session folder, instead of relying on the shared system
  Downloads folder (which is what made concurrent runs unsafe before).
- Added a "Download report" button so the finished file can be downloaded
  straight from the browser once the server has generated it, instead of
  needing to be located on the server's disk.
- Login credentials are now read from environment variables instead of
  being hardcoded in the source (set AKSJEEIERE_EMAIL / AKSJEEIERE_PASSWORD,
  e.g. via a .env file).
- All comments and interface text translated to English.
"""

# ============================================================
# 1. IMPORTS
# ============================================================
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select
from selenium.common.exceptions import NoSuchElementException, WebDriverException, TimeoutException
from webdriver_manager.chrome import ChromeDriverManager


import os
import re
import sys
import time
import random
import shutil
import subprocess
import glob
import uuid

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment

import streamlit as st

# Optional: load AKSJEEIERE_EMAIL / AKSJEEIERE_PASSWORD from a local .env
# file if python-dotenv is installed. Safe to leave uninstalled - the
# program still works if credentials are set directly as environment
# variables (e.g. on the server) instead of via a .env file.
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

AKSJEEIERE_EMAIL = os.environ.get("AKSJEEIERE_EMAIL", "")
AKSJEEIERE_PASSWORD = os.environ.get("AKSJEEIERE_PASSWORD", "")


# ============================================================
# 2. PER-SESSION SETUP (this is what makes it safe for several
#    people to use the app on the server at the same time)
# ============================================================
# Every browser tab/user that opens the app gets its own session_id,
# generated once and stored in Streamlit's session state (so it stays the
# same across reruns for that user, but is different for every other
# user). All working files for a given run - downloaded xlsx files and the
# finished report - live inside a folder named after this ID
# (sessions/<session_id>/), so two people running the app at the same time
# never read, overwrite, or delete each other's files.
if "session_id" not in st.session_state:
    st.session_state.session_id = uuid.uuid4().hex[:8]

session_id = st.session_state.session_id
session_folder = os.path.join(os.getcwd(), 'sessions', session_id)
os.makedirs(session_folder, exist_ok=True)


# ============================================================
# 3. SAFE NUMBER-CONVERSION HELPER FUNCTIONS
# ============================================================
def safe_int(value, default=0):
    """
    Safely converts a value to int.
    If the value is empty, whitespace, NaN, or non-numeric text, returns
    `default` instead of crashing (this is exactly where the program used
    to break before).
    """
    if value is None:
        return default
    try:
        # pd.isna catches NaN/None; str(...).strip() removes whitespace like " "
        if pd.isna(value):
            return default
    except (TypeError, ValueError):
        pass

    text_value = str(value).strip()
    if text_value == "" or text_value.lower() == "nan":
        return default

    try:
        return int(float(text_value))  # float() helps with values like "123.0"
    except (ValueError, TypeError):
        print(f"⚠️ Could not convert value '{value}' to a number, using {default}")
        return default


def sanitize_filename(name):
    """Replaces characters that aren't allowed in Windows/Mac/Linux filenames with '_'."""
    return re.sub(r'[\\/*?:"<>|]', '_', name)


def normalize_company_name(name):
    """
    Normalizes a company name for comparison, so that formatting differences
    don't cause a mismatch - e.g. "Heidenreich Holding A/S",
    "heidenreich holding as" and "Heidenreich Holding A.S." are all treated
    as the same company. Removes slashes and dots, lowercases, and
    collapses extra whitespace.
    """
    name = str(name).strip().lower()
    name = name.replace('/', '').replace('.', '')
    name = re.sub(r'\s+', ' ', name).strip()
    return name


# ============================================================
# 4. SELENIUM: DOWNLOADING DATA FROM AKSJEEIERE.NO
# ============================================================
import os

def build_chrome_driver(download_dir, headless=True):
    """
    Creates a Chrome WebDriver configured to download files directly into
    `download_dir` (this session's own folder).
    """
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")

    options.add_experimental_option("prefs", {
        "download.default_directory": download_dir,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "safebrowsing.enabled": True,
    })

    # Перевіряємо, чи працюємо у середовищі Streamlit Cloud (Linux)
    if os.path.exists("/usr/bin/chromium"):
        options.binary_location = "/usr/bin/chromium"
        service = Service("/usr/bin/chromedriver")
    else:
        # Для локальної розробки (на власному ПК)
        from webdriver_manager.chrome import ChromeDriverManager
        service = Service(ChromeDriverManager().install())

    driver = webdriver.Chrome(service=service, options=options)
    return driver

def daselenium(company_name_input, session_folder, headless=True):
    """
    Logs into aksjeeiere.no, searches for the company, and downloads xlsx
    files with the shareholder list for each year from 2014 to
    (current year - 1). Files download directly into this session's own
    folder (session_folder), so concurrent runs from different users never
    collide.
    """
    driver = None
    try:
        driver = build_chrome_driver(download_dir=session_folder, headless=headless)
        print("✅ ChromeDriver started successfully")

        driver.get("https://www.aksjeeiere.no/")
        time.sleep(random.uniform(1, 3))

        # --- Login ---
       # Чекаємо до 10 секунд, поки з'явиться посилання на логін
    try:
       log_in = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.XPATH, "//a[contains(@href, 'login') or contains(text(), 'Logg inn') or contains(text(), 'Log in')]"))
    )
          log_in.click()
    except Exception as e:
          print("⚠️ unable to find lohin button:", e)
          raise

        time.sleep(0.5)

        try:
            email_input = driver.find_element(By.XPATH, "//input[@id='email']")
            email_input.send_keys(AKSJEEIERE_EMAIL)

            time.sleep(0.5)

            password_input = driver.find_element(By.XPATH, "//input[@id='password']")
            password_input.send_keys(AKSJEEIERE_PASSWORD)

            time.sleep(0.5)

            login_input = driver.find_element(By.XPATH, "//input[@type='submit']")
            login_input.click()
        except NoSuchElementException:
            print("⚠️ Could not find the login/password fields - check the login page structure")
            raise

        time.sleep(1)

        # --- Search for the company ---
        try:
            search_bar = driver.find_element(By.ID, "q")
            search_bar.send_keys(company_name_input)
        except NoSuchElementException:
            print("⚠️ Could not find the company search field ('q')")
            raise

        time.sleep(0.5)

        # Years to download data for: from 2014 to (current year - 1)
        current_year = time.localtime().tm_year
        downloaded_years = []
        skipped_years = []

        for year in range(2014, current_year):
            try:
                dropdown = Select(driver.find_element(By.ID, "year"))
                dropdown.select_by_visible_text(str(year))
            except NoSuchElementException:
                print(f"⚠️ Could not find the year dropdown for {year}, skipping")
                skipped_years.append(year)
                continue

            time.sleep(random.uniform(1, 3))

            try:
                download_button = driver.find_element(
                    By.XPATH,
                    "//button[@class= 'sm:w-auto w-full sm:relative sm:flex items-center space-x-2 sm:rounded-r-md bg-sky-600 px-4 py-3 sm:px-4 sm:py-2 text-sm focus-visible:outline focus-visible:outline-2 focus-visible:outline-offset-2 bg-sky-600 text-white shadow-sm hover:bg-sky-500 focus-visible:outline-sky-600']"
                )
                download_button.click()
            except NoSuchElementException:
                print(f"⚠️ Could not find the search button for {year}, skipping")
                skipped_years.append(year)
                continue

            time.sleep(1)

            try:
                # Link to download the result (xlsx)
                link = driver.find_element(By.LINK_TEXT, "Last ned søkeresultat (xlsx)")
                link.click()
                time.sleep(1.5)  # give the file time to download
                downloaded_years.append(year)
            except NoSuchElementException:
                # No data for this year - this is normal, not an error
                print(f"ℹ️ No data for {year}, skipping")
                skipped_years.append(year)
                continue

        print(f"✅ Years downloaded: {downloaded_years}")
        if skipped_years:
            print(f"ℹ️ Years skipped (no data or elements on the page): {skipped_years}")

    except WebDriverException as e:
        print(f"❌ Selenium/ChromeDriver error: {e}")
        raise
    except Exception as e:
        print(f"❌ Unexpected error while downloading data: {e}")
        raise

    finally:
        if driver is not None:
            try:
                logout = driver.find_element(By.XPATH, "//a[@rel='nofollow']")
                logout.click()
            except Exception as e:
                print("ℹ️ Logout link not found or not needed:", e)
            driver.quit()
        else:
            print("ℹ️ Driver was never initialized, logout skipped.")


# ============================================================
# 5. EXCEL PROCESSING: BUILDING THE FINAL REPORT
# ============================================================
def write_data(company_name_input, session_folder):
    """
    Reads every xlsx file from this session's folder, merges them into the
    Template(original).xlsx template, calculates year-over-year changes in
    share count, and saves the final report (also inside the session
    folder, so concurrent runs never overwrite each other's output).
    Returns the path to the finished file.
    """

    # --- 5.1 Find the shareholder files and the template ---
    excel_files = glob.glob(os.path.join(session_folder, "*.xlsx"))
    print(f"ℹ️ Found {len(excel_files)} xlsx files in this session's folder")

    if not excel_files:
        # Previously the program would carry on and crash on empty data -
        # now it stops right away instead.
        print("❌ No xlsx files in this session's folder - nothing to process")
        return 'This company has no shareholders'

    template_path = os.path.join(os.getcwd(), 'Template(original).xlsx')
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"Template file not found: {template_path}. "
            "Make sure 'Template(original).xlsx' is in the same folder as the program."
        )

    book = load_workbook(template_path)
    print('✅ Excel template loaded')

    sheet_name = 'Private & Selskap 2014 - 2024'
    if sheet_name not in book.sheetnames:
        raise ValueError(f"Template is missing sheet '{sheet_name}'. Available sheets: {book.sheetnames}")

    template_df = pd.read_excel(template_path, sheet_name=sheet_name)
    original_columns = template_df.columns.tolist()

    # Mapping from aksjeeiere.no column names -> template column names
    column_mapping = {
        'Aksjonær': 'SHAREHOLDER',
        'Født/org.nr.': 'SEC.NR. / ORG.NR.',
        'Postnr./sted': 'ZIP CODE / CITY',
        'Selskap': 'COMPANY',
        'År': 'YEAR',
        'Aksjer': 'TOTAL NR. SHARES',
    }

    company_name = normalize_company_name(company_name_input)

    # --- 5.2 Read each file, filter by company, and append to the combined table ---
    for idx, aksjeeiere_path in enumerate(excel_files):
        print(f"ℹ️ Processing file [{idx + 1}/{len(excel_files)}]: {aksjeeiere_path}")
        aksjeeiere_df = pd.read_excel(aksjeeiere_path)

        if 'Selskap' not in aksjeeiere_df.columns:
            print(f"⚠️ File {aksjeeiere_path} has no 'Selskap' column, skipping file")
            continue

        available_companies = aksjeeiere_df['Selskap'].unique()
        print(f"   Companies in file: {list(available_companies)}")
        print(f"   Looking for: {company_name}")

        aksjeeiere_df = aksjeeiere_df[aksjeeiere_df['Selskap'].apply(normalize_company_name) == company_name]
        print(f"   Rows found: {aksjeeiere_df.shape[0]}")

        if aksjeeiere_df.empty:
            print(f"   ℹ️ No data for this company in this file, skipping")
            continue

        aksjeeiere_df = aksjeeiere_df.sort_values(by='Aksjer', ascending=False)

        # If a company name wasn't explicitly entered - take it from the first matching row
        if not company_name_input:
            company_name = str(aksjeeiere_df['Selskap'].iloc[0]).replace(' ', '_').replace('/', '_')

        # Build a temporary table with columns in the template's format
        temp_df = pd.DataFrame()
        for aksjeeiere_col, template_col in column_mapping.items():
            if aksjeeiere_col in aksjeeiere_df.columns and template_col in template_df.columns:
                temp_df[template_col] = aksjeeiere_df[aksjeeiere_col]
            else:
                print(f"   ⚠️ Column '{aksjeeiere_col}' missing from file or template - skipping this mapping")

        # YEAR needs to be numeric. astype(int) used to crash on non-numeric
        # values - now we "softly" convert first and drop invalid rows.
        template_df['YEAR'] = pd.to_numeric(template_df['YEAR'], errors='coerce')
        rows_before = len(template_df)
        template_df = template_df.dropna(subset=['YEAR'])
        rows_dropped = rows_before - len(template_df)
        if rows_dropped:
            print(f"   ⚠️ Removed {rows_dropped} rows with non-numeric/empty YEAR")
        template_df['YEAR'] = template_df['YEAR'].astype(int)

        template_df = pd.concat([template_df, temp_df], ignore_index=True)
        all_columns = original_columns + [col for col in template_df.columns if col not in original_columns]
        template_df = template_df.reindex(columns=all_columns, fill_value="")

    print(f"ℹ️ Final table size: {template_df.shape}")

    if template_df.empty:
        print(f"❌ No shareholders found for company '{company_name}'")
        return 'This company has no shareholders'

    # --- 5.3 Calculate year-over-year changes in share count ---
    all_changes_dict = {}
    previous_year_shares = {}
    previous_year_keys = set()

    def normalize_name(name, sec_nr):
        """For private individuals (short SEC.NR.), sorts the words in the name to avoid duplicates caused by word order."""
        name = str(name)
        if len(sec_nr) <= 5:
            return " ".join(sorted(name.split()))
        return name

    latest_year = template_df["YEAR"].max() + 1
    earliest_year = template_df["YEAR"].min()
    years = list(range(2014, latest_year))

    for year in years:
        df_current_year = template_df[template_df['YEAR'] == year]
        current_year_keys = set()

        for row_idx, row in df_current_year.iterrows():
            comment = ""
            shareholder = row['SHAREHOLDER']

            # MAIN BUG FIX: this used to be int(sec_nr), which crashed with
            # "invalid literal for int() with base 10: ' '" if the cell
            # held a space instead of a number/NaN.
            # safe_int() correctly handles empty values, whitespace, NaN and text.
            sec_nr = str(safe_int(row['SEC.NR. / ORG.NR.'], default=0))

            # Share count could also be non-numeric (text, whitespace) - handle safely
            total_shares = safe_int(row['TOTAL NR. SHARES'], default=0)

            normalized_shareholder = normalize_name(shareholder, sec_nr)
            key = (normalized_shareholder, sec_nr)
            current_year_keys.add(key)

            if key not in all_changes_dict and earliest_year is not None:
                all_changes_dict[key] = []
                if year > earliest_year:
                    comment = f"New {total_shares} shares in {year}"
                    all_changes_dict[key].append(comment)
                print(f"➕ Added new shareholder {key} to the dictionary: {all_changes_dict[key]}")

            if key in previous_year_shares:
                last_year_shares = previous_year_shares[key]
                if total_shares > last_year_shares:
                    comment = f"Bought {total_shares - last_year_shares} shares in {year}"
                    all_changes_dict[key].append(comment)
                elif total_shares < last_year_shares:
                    comment = f"Sold {last_year_shares - total_shares} shares in {year}"
                    all_changes_dict[key].append(comment)

            if comment:
                template_df.loc[template_df.index == row_idx, 'CHANGES 2014-2024'] = comment
                print(f"📝 Comment for {key}: {comment}")

            previous_year_shares[key] = total_shares

        # Shareholders who disappeared (sold all shares) - add a comment in the previous year
        if year > earliest_year:
            missing_keys = previous_year_keys - current_year_keys
            for missing_key in missing_keys:
                comment = f"Sold all shares in {year}"
                all_changes_dict[missing_key].append(comment)

                normalized_name, sec_nr = missing_key
                mask = (
                    (template_df['SHAREHOLDER'].apply(lambda x: normalize_name(x, sec_nr)) == normalized_name) &
                    (template_df['SEC.NR. / ORG.NR.'].apply(lambda v: str(safe_int(v, 0))) == sec_nr)
                )
                last_year_rows = template_df[mask & (template_df['YEAR'] == year - 1)]

                if not last_year_rows.empty:
                    last_idx = last_year_rows.index[-1]
                    existing_comment = template_df.at[last_idx, 'CHANGES 2014-2024']
                    if pd.isna(existing_comment):
                        template_df.at[last_idx, 'CHANGES 2014-2024'] = comment
                    else:
                        template_df.at[last_idx, 'CHANGES 2014-2024'] += f"; {comment}"
                    print(f"⚠️ Shareholder {missing_key} disappeared - comment added to year {year - 1}: {comment}")

        previous_year_keys = current_year_keys

    print('✅ CHANGES 2014-2024 column filled in')

    # --- 5.4 Split into private individuals vs companies (by SEC.NR. length) ---
    df_latest_year = template_df[template_df["YEAR"] == latest_year - 1].copy()
    df_latest_year.loc[:, 'SEC.NR. / ORG.NR.'] = df_latest_year['SEC.NR. / ORG.NR.'].apply(lambda v: str(safe_int(v, 0)))

    df_private = df_latest_year[df_latest_year["SEC.NR. / ORG.NR."].str.len() <= 5].copy()
    df_selskap = df_latest_year[df_latest_year["SEC.NR. / ORG.NR."].str.len() > 5].copy()

    if not df_private.empty:
        df_private['normalized_shareholder'] = df_private.apply(
            lambda row: normalize_name(row['SHAREHOLDER'], row['SEC.NR. / ORG.NR.']), axis=1
        )
    else:
        print("ℹ️ df_private is empty - skipping normalized_shareholder")

    # --- 5.5 Apply the accumulated comments to both tables ---
    for (shareholder, sec_nr), comments in all_changes_dict.items():
        normalized_key = normalize_name(shareholder, sec_nr)

        if not df_private.empty and 'normalized_shareholder' in df_private.columns:
            mask = (df_private['normalized_shareholder'] == normalized_key) & (df_private['SEC.NR. / ORG.NR.'] == sec_nr)
            if mask.any():
                df_private.loc[mask, 'CHANGES 2014-2024'] = "\n".join(comments)

        if not df_selskap.empty:
            mask2 = (df_selskap["SHAREHOLDER"] == shareholder) & (df_selskap['SEC.NR. / ORG.NR.'] == sec_nr)
            if mask2.any():
                df_selskap.loc[mask2, 'CHANGES 2014-2024'] = "\n".join(comments)

    template_df = template_df.sort_values(by="YEAR", ascending=False)

    if not df_private.empty and 'normalized_shareholder' in df_private.columns:
        df_private = df_private.drop('normalized_shareholder', axis=1)

    # --- 5.6 Save the final file (inside this session's own folder) ---
    safe_company_name = sanitize_filename(company_name)
    output_path = os.path.join(session_folder, f'{safe_company_name} {(latest_year) - 1}.xlsx')
    book.save(output_path)

    with pd.ExcelWriter(output_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
        df_private.to_excel(writer, sheet_name="Private 2024", index=False)
        df_selskap.to_excel(writer, sheet_name="Selskap 2024", index=False)
        template_df.to_excel(writer, sheet_name="Private & Selskap 2014 - 2024", index=False)

    _apply_wrap_text(output_path, ["Private 2024", "Selskap 2024", "Private & Selskap 2014 - 2024"])

    # --- 5.7 Per-year summary on the "Analyse" sheet ---
    wb = load_workbook(output_path)
    df_private_selskap = pd.read_excel(output_path, sheet_name="Private & Selskap 2014 - 2024")

    # Guard against the same bug: clean up non-numeric values before grouping/summing
    df_private_selskap['YEAR'] = pd.to_numeric(df_private_selskap['YEAR'], errors='coerce')
    df_private_selskap['TOTAL NR. SHARES'] = pd.to_numeric(df_private_selskap['TOTAL NR. SHARES'], errors='coerce')
    df_private_selskap = df_private_selskap.dropna(subset=['YEAR', 'TOTAL NR. SHARES'])

    year_summary = df_private_selskap.groupby('YEAR')['TOTAL NR. SHARES'].sum().reset_index()

    if 'Analyse' not in wb.sheetnames:
        print("⚠️ Template has no 'Analyse' sheet - skipping the per-year summary")
    else:
        ws_analyse = wb["Analyse"]
        summary_text = "\n".join(
            f"{int(row['YEAR'])} : {int(row['TOTAL NR. SHARES'])}" for _, row in year_summary.iterrows()
        )
        ws_analyse['U3'] = summary_text
        ws_analyse['U3'].alignment = Alignment(wrap_text=True)

    wb.save(output_path)

    print(f"✅ Finished file: {output_path}")
    return output_path


def _apply_wrap_text(output_path, sheet_names, column_header="CHANGES 2014-2024"):
    """Enables text wrapping for the column_header column on the given sheets."""
    wb = load_workbook(output_path)
    for sheet_name in sheet_names:
        if sheet_name not in wb.sheetnames:
            print(f"⚠️ Sheet '{sheet_name}' not found, skipping wrap text")
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=1):  # headers only
            for cell in row:
                if cell.value == column_header:
                    col_letter = cell.column_letter
                    for data_cell in ws[col_letter]:
                        data_cell.alignment = Alignment(wrap_text=True)
                    break
    wb.save(output_path)


# ============================================================
# 6. CLEANING UP TEMPORARY FILES
# ============================================================
# NOTE: this only ever deletes THIS session's own folder
# (sessions/<session_id>/), never a folder shared with other users - that
# isolation is what makes it safe to call even while other people are
# using the app on the server at the same time.
def delete_session_folder(session_folder, max_attempts=5, delay_seconds=1):
    """
    Deletes this session's entire working folder (downloaded input files +
    finished report). Retries a few times with a short pause if deletion
    fails with a permission error - this happens occasionally when the
    folder is inside a OneDrive-synced location and OneDrive briefly holds
    a file while uploading it to the cloud.
    """
    if not os.path.exists(session_folder):
        print(f"ℹ️ Session folder {session_folder} does not exist, nothing to clean up")
        return

    for attempt in range(1, max_attempts + 1):
        try:
            shutil.rmtree(session_folder)
            print(f"ℹ️ Cleaned up session folder {session_folder}")
            return
        except PermissionError as e:
            print(f"⚠️ Attempt {attempt}/{max_attempts} to delete session folder failed ({e}), retrying...")
            time.sleep(delay_seconds)

    print(f"⚠️ Could not delete session folder {session_folder} after {max_attempts} attempts - leaving it in place")


# ============================================================
# 7. STREAMLIT UI
# ============================================================
st.title("Excel template maker")

company_name_input = st.text_input("Enter Company Name:")
include_phone_numbers = st.checkbox("Include Telephone Numbers")

if st.button("Search and Download"):
    if not company_name_input:
        st.warning("Please enter a company name before starting.")
    elif not AKSJEEIERE_EMAIL or not AKSJEEIERE_PASSWORD:
        st.error(
            "Missing AKSJEEIERE_EMAIL / AKSJEEIERE_PASSWORD environment variables. "
            "Set them (e.g. in a .env file) before running."
        )
    else:
        try:
            # Step 1: download data from aksjeeiere.no via Selenium
            with st.spinner("Downloading data from aksjeeiere.no..."):
                daselenium(company_name_input, session_folder)

            # Step 2: process the data and assemble the final Excel report
            with st.spinner("Processing data and building the report..."):
                _file_pathexcel = write_data(company_name_input, session_folder)
            print(f"✅ Excel file created: {_file_pathexcel}")

            if _file_pathexcel == 'This company has no shareholders':
                st.warning("No shareholders found for this company. Please check the company name and try again.")
                print("ℹ️ No shareholders found - skipping telephone extraction")
            else:
                # Step 3: (optional) add phone numbers
                if include_phone_numbers:
                    telephone_script = os.path.join(os.getcwd(), "telephone.py")
                    if not os.path.exists(telephone_script):
                        st.error("telephone.py not found - skipping phone number lookup.")
                        print(f"❌ telephone.py not found at {telephone_script}")
                    elif not os.path.exists(_file_pathexcel):
                        st.error(f"Excel file not found: {_file_pathexcel}")
                        print(f"❌ ERROR: Excel file not found: {_file_pathexcel}")
                    else:
                        with st.spinner("Adding phone numbers..."):
                            result = subprocess.run(
                                [sys.executable, telephone_script, _file_pathexcel],
                                capture_output=True, text=True, encoding="utf-8", errors="replace"
                            )
                        if result.returncode != 0:
                            st.error(f"Telephone extraction failed: {result.stderr}")
                            print(f"❌ Telephone script error: {result.stderr}")
                        else:
                            print("✅ Telephone extraction completed successfully")
                            # telephone.py saves the result under the same name
                            # plus a " (T)" suffix - point the download button
                            # at that file instead of the original.
                            base_name = os.path.splitext(os.path.basename(_file_pathexcel))[0]
                            renamed_matches = sorted(
                                glob.glob(os.path.join(session_folder, f"{base_name} (T)*.xlsx")),
                                key=os.path.getmtime,
                                reverse=True,
                            )
                            if renamed_matches:
                                _file_pathexcel = renamed_matches[0]

                st.success("Processing completed successfully!")

                # Step 4: offer the finished file for download directly in the
                # browser. st.download_button reads the file's bytes into the
                # page right away, so it's safe to clean up the session folder
                # afterwards (below) even before the user clicks it.
                with open(_file_pathexcel, "rb") as f:
                    st.download_button(
                        label="📥 Download report",
                        data=f,
                        file_name=os.path.basename(_file_pathexcel),
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )

        except Exception as e:
            st.error(f"An error occurred: {str(e)}")
            print(f"❌ Error in main processing: {e}")

        finally:
            # Step 5: clean up this session's temporary files (downloaded
            # xlsx files + finished report), regardless of success or
            # failure. Only ever touches this session's own folder.
            try:
                delete_session_folder(session_folder)
            except Exception as cleanup_error:
                print(f"⚠️ Could not clean up temporary files: {cleanup_error}")
